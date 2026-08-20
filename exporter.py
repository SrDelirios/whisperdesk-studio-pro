import os
import re
import io
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
)
from reportlab.pdfgen import canvas

class NumberedCanvas(canvas.Canvas):
    """Canvas de ReportLab que calcula el total de páginas y agrega pie de página ejecutivo."""
    def __init__(self, *args, **kwargs):
        super(NumberedCanvas, self).__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super(NumberedCanvas, self).showPage()
        super(NumberedCanvas, self).save()

    def draw_page_decorations(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748B"))
        
        self.setStrokeColor(colors.HexColor("#E2E8F0"))
        self.setLineWidth(0.5)
        self.line(40, 36, letter[0] - 40, 36)
        
        self.drawString(40, 24, "WhisperDesk Studio Pro — Acta Oficial de Sesión")
        self.drawRightString(letter[0] - 40, 24, f"Página {self._pageNumber} de {page_count}")
        self.restoreState()


class MeetingMinutesExporter:
    """Exportador profesional de Actas de Reunión a DOCX (Word), PDF y Markdown."""

    @staticmethod
    def _parse_markdown_blocks(markdown_text: str):
        """Parsea el texto Markdown del acta en bloques estructurados para los generadores."""
        lines = [l.rstrip() for l in markdown_text.split('\n')]
        blocks = []
        current_list = None

        for line in lines:
            trimmed = line.strip()
            if not trimmed:
                if current_list:
                    blocks.append(current_list)
                    current_list = None
                continue

            if trimmed.startswith('# '):
                if current_list:
                    blocks.append(current_list)
                    current_list = None
                blocks.append({"type": "h1", "text": trimmed[2:].strip()})
            elif trimmed.startswith('## '):
                if current_list:
                    blocks.append(current_list)
                    current_list = None
                blocks.append({"type": "h2", "text": trimmed[3:].strip()})
            elif trimmed.startswith('### '):
                if current_list:
                    blocks.append(current_list)
                    current_list = None
                blocks.append({"type": "h3", "text": trimmed[4:].strip()})
            elif trimmed.startswith('#### '):
                if current_list:
                    blocks.append(current_list)
                    current_list = None
                blocks.append({"type": "h4", "text": trimmed[5:].strip()})
            elif trimmed.startswith('---') or trimmed.startswith('***') or trimmed.startswith('==='):
                if current_list:
                    blocks.append(current_list)
                    current_list = None
                blocks.append({"type": "hr"})
            elif trimmed.startswith('- [ ] ') or trimmed.startswith('- [x] ') or trimmed.startswith('- [X] '):
                is_checked = trimmed.startswith('- [x] ') or trimmed.startswith('- [X] ')
                item_text = trimmed[6:].strip()
                if not current_list or current_list["type"] != "checklist":
                    if current_list:
                        blocks.append(current_list)
                    current_list = {"type": "checklist", "items": []}
                current_list["items"].append({"checked": is_checked, "text": item_text})
            elif trimmed.startswith('- ') or trimmed.startswith('* '):
                item_text = trimmed[2:].strip()
                if not current_list or current_list["type"] != "bullet":
                    if current_list:
                        blocks.append(current_list)
                    current_list = {"type": "bullet", "items": []}
                current_list["items"].append(item_text)
            elif re.match(r'^\d+\.\s+', trimmed):
                item_text = re.sub(r'^\d+\.\s+', '', trimmed)
                if not current_list or current_list["type"] != "numbered":
                    if current_list:
                        blocks.append(current_list)
                    current_list = {"type": "numbered", "items": []}
                current_list["items"].append(item_text)
            elif trimmed.startswith('> '):
                if current_list:
                    blocks.append(current_list)
                    current_list = None
                blocks.append({"type": "quote", "text": trimmed[2:].strip()})
            else:
                if current_list:
                    blocks.append(current_list)
                    current_list = None
                blocks.append({"type": "p", "text": trimmed})

        if current_list:
            blocks.append(current_list)

        return blocks

    @staticmethod
    def generate_docx(markdown_text: str, title: str = "Acta Oficial de Sesión") -> io.BytesIO:
        """Genera un documento Word (.docx) con diseño ejecutivo de alta calidad."""
        doc = Document()

        for section in doc.sections:
            section.top_margin = Inches(0.8)
            section.bottom_margin = Inches(0.8)
            section.left_margin = Inches(0.85)
            section.right_margin = Inches(0.85)

        COLOR_PRIMARY = RGBColor(30, 41, 59)      # Slate 800
        COLOR_ACCENT = RGBColor(99, 102, 241)     # Indigo 500
        COLOR_TEXT = RGBColor(51, 65, 85)         # Slate 700
        COLOR_MUTED = RGBColor(100, 116, 139)     # Slate 500

        style_normal = doc.styles['Normal']
        style_normal.font.name = 'Calibri'
        style_normal.font.size = Pt(10.5)
        style_normal.font.color.rgb = COLOR_TEXT

        blocks = MeetingMinutesExporter._parse_markdown_blocks(markdown_text)

        def add_formatted_runs(paragraph, text, base_bold=False, base_color=COLOR_TEXT, base_size=Pt(10.5)):
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if not part:
                    continue
                run = paragraph.add_run()
                run.font.name = 'Calibri'
                run.font.size = base_size
                run.font.color.rgb = base_color
                if part.startswith('**') and part.endswith('**') and len(part) >= 4:
                    run.text = part[2:-2]
                    run.bold = True
                else:
                    run.text = part
                    run.bold = base_bold

        for block in blocks:
            b_type = block["type"]

            if b_type == "h1":
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(14)
                p.paragraph_format.space_after = Pt(4)
                add_formatted_runs(p, block["text"], base_bold=True, base_color=COLOR_ACCENT, base_size=Pt(18))
                
            elif b_type == "h2":
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(12)
                p.paragraph_format.space_after = Pt(4)
                add_formatted_runs(p, block["text"], base_bold=True, base_color=COLOR_PRIMARY, base_size=Pt(14))

            elif b_type == "h3":
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after = Pt(3)
                add_formatted_runs(p, block["text"], base_bold=True, base_color=COLOR_PRIMARY, base_size=Pt(12))

            elif b_type == "h4":
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(2)
                add_formatted_runs(p, block["text"], base_bold=True, base_color=COLOR_PRIMARY, base_size=Pt(11))

            elif b_type == "hr":
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(6)
                pBdr = OxmlElement('w:pBdr')
                bottom = OxmlElement('w:bottom')
                bottom.set(qn('w:val'), 'single')
                bottom.set(qn('w:sz'), '6')
                bottom.set(qn('w:space'), '1')
                bottom.set(qn('w:color'), 'CBD5E1')
                pBdr.append(bottom)
                p._p.get_or_add_pPr().append(pBdr)

            elif b_type == "p":
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(5)
                p.paragraph_format.line_spacing = 1.15
                add_formatted_runs(p, block["text"])

            elif b_type == "quote":
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.3)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.line_spacing = 1.15
                add_formatted_runs(p, block["text"], base_color=COLOR_MUTED)

            elif b_type == "bullet":
                for item in block["items"]:
                    p = doc.add_paragraph(style='List Bullet')
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.line_spacing = 1.15
                    add_formatted_runs(p, item)

            elif b_type == "numbered":
                for idx, item in enumerate(block["items"], 1):
                    p = doc.add_paragraph(style='List Number')
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.line_spacing = 1.15
                    add_formatted_runs(p, item)

            elif b_type == "checklist":
                for item in block["items"]:
                    p = doc.add_paragraph()
                    p.paragraph_format.left_indent = Inches(0.2)
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.space_after = Pt(3)
                    
                    check_symbol = "☑ " if item["checked"] else "☐ "
                    box_run = p.add_run(check_symbol)
                    box_run.font.name = 'Arial'
                    box_run.font.size = Pt(11)
                    box_run.bold = True
                    box_run.font.color.rgb = COLOR_ACCENT if item["checked"] else COLOR_MUTED

                    add_formatted_runs(p, item["text"])

        output = io.BytesIO()
        doc.save(output)
        output.seek(0)
        return output

    @staticmethod
    def generate_pdf(markdown_text: str, title: str = "Acta Oficial de Sesión") -> io.BytesIO:
        """Genera un archivo PDF ejecutivo con diseño corporativo y paginación."""
        output = io.BytesIO()
        doc = SimpleDocTemplate(
            output,
            pagesize=letter,
            leftMargin=42,
            rightMargin=42,
            topMargin=45,
            bottomMargin=45
        )

        styles = getSampleStyleSheet()
        PRIMARY_COLOR = colors.HexColor("#0F172A")
        ACCENT_COLOR = colors.HexColor("#4F46E5")
        TEXT_COLOR = colors.HexColor("#334155")
        MUTED_COLOR = colors.HexColor("#64748B")
        BORDER_COLOR = colors.HexColor("#CBD5E1")

        style_h1 = ParagraphStyle(
            'ActaH1',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=16,
            leading=20,
            textColor=ACCENT_COLOR,
            spaceBefore=10,
            spaceAfter=6
        )

        style_h2 = ParagraphStyle(
            'ActaH2',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=13,
            leading=16,
            textColor=PRIMARY_COLOR,
            spaceBefore=12,
            spaceAfter=5
        )

        style_h3 = ParagraphStyle(
            'ActaH3',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=11,
            leading=14,
            textColor=PRIMARY_COLOR,
            spaceBefore=9,
            spaceAfter=4
        )

        style_body = ParagraphStyle(
            'ActaBody',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9.5,
            leading=13.5,
            textColor=TEXT_COLOR,
            spaceBefore=2,
            spaceAfter=4
        )

        style_bullet = ParagraphStyle(
            'ActaBullet',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9.5,
            leading=13.5,
            textColor=TEXT_COLOR,
            leftIndent=15,
            spaceBefore=1,
            spaceAfter=3
        )

        style_quote = ParagraphStyle(
            'ActaQuote',
            parent=styles['Normal'],
            fontName='Helvetica-Oblique',
            fontSize=9,
            leading=13,
            textColor=MUTED_COLOR,
            leftIndent=18,
            spaceBefore=3,
            spaceAfter=5
        )

        blocks = MeetingMinutesExporter._parse_markdown_blocks(markdown_text)
        story = []

        def md_to_reportlab_html(text):
            escaped = text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
            escaped = re.sub(r'\*\*(.*?)\*\*', r'<b>\1</b>', escaped)
            escaped = re.sub(r'\*(.*?)\*', r'<i>\1</i>', escaped)
            return escaped

        for block in blocks:
            b_type = block["type"]

            if b_type == "h1":
                story.append(Paragraph(md_to_reportlab_html(block["text"]), style_h1))
            elif b_type == "h2":
                story.append(Paragraph(md_to_reportlab_html(block["text"]), style_h2))
            elif b_type == "h3":
                story.append(Paragraph(md_to_reportlab_html(block["text"]), style_h3))
            elif b_type == "h4":
                story.append(Paragraph(md_to_reportlab_html(block["text"]), style_h3))
            elif b_type == "hr":
                story.append(Spacer(1, 3))
                story.append(HRFlowable(width="100%", thickness=0.6, color=BORDER_COLOR, spaceBefore=4, spaceAfter=6))
            elif b_type == "p":
                story.append(Paragraph(md_to_reportlab_html(block["text"]), style_body))
            elif b_type == "quote":
                story.append(Paragraph(md_to_reportlab_html(block["text"]), style_quote))
            elif b_type == "bullet":
                for item in block["items"]:
                    bullet_text = f"&bull;&nbsp;&nbsp;{md_to_reportlab_html(item)}"
                    story.append(Paragraph(bullet_text, style_bullet))
            elif b_type == "numbered":
                for idx, item in enumerate(block["items"], 1):
                    num_text = f"<b>{idx}.</b>&nbsp;&nbsp;{md_to_reportlab_html(item)}"
                    story.append(Paragraph(num_text, style_bullet))
            elif b_type == "checklist":
                for item in block["items"]:
                    box = "<font color='#4F46E5'><b>[✓]</b></font>" if item["checked"] else "<font color='#94A3B8'>[ ]</font>"
                    check_text = f"{box}&nbsp;&nbsp;{md_to_reportlab_html(item['text'])}"
                    story.append(Paragraph(check_text, style_bullet))

        doc.build(story, canvasmaker=NumberedCanvas)
        output.seek(0)
        return output
